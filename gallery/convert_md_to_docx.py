import sys, re, os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert(md_path, docx_path):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.first_line_indent = Cm(1.25)

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code = False
    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith('```'):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            i += 1
            continue

        img_match = re.match(r'^!\[(.+?)\]\((.+?)\)$', line)
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            doc.add_paragraph()
            if os.path.exists(img_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p_img.add_run()
                run.add_picture(img_path, width=Inches(4.5))
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(f'[Image not found: {img_path}]').bold = True
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.add_run(caption).italic = False
            p_cap.paragraph_format.first_line_indent = Cm(0)
            doc.add_paragraph()
            i += 1
            continue

        if line == '':
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith('# '):
            text = line[2:]
            doc.add_paragraph()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text.upper())
            run.bold = True
            run.font.size = Pt(16)
            p.paragraph_format.first_line_indent = Cm(0)
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith('## '):
            text = line[3:]
            p = doc.add_paragraph()
            run = p.add_run(text.upper())
            run.bold = True
            run.font.size = Pt(14)
            p.paragraph_format.first_line_indent = Cm(1.25)
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith('### '):
            text = line[4:]
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(14)
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith('- '):
            p = doc.add_paragraph()
            p.add_run(line[2:])
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(1.5)
            i += 1
            continue

        if line.startswith('1. '):
            p = doc.add_paragraph()
            p.add_run(line)
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(1.5)
            i += 1
            continue

        if line.startswith('| ') and '|' in line:
            # Table row - collect all rows
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip())
                i += 1
            if len(rows) >= 2:
                # Parse header
                headers = [h.strip() for h in rows[0].split('|') if h.strip()]
                # Parse data rows (skip separator row)
                data_rows = []
                for r in rows[2:]:
                    cells = [c.strip() for c in r.split('|') if c.strip()]
                    if cells:
                        data_rows.append(cells)
                if data_rows:
                    table = doc.add_table(rows=1 + len(data_rows), cols=len(headers))
                    table.style = 'Table Grid'
                    for j, h in enumerate(headers):
                        cell = table.rows[0].cells[j]
                        cell.text = h
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in p.runs:
                                run.bold = True
                                run.font.name = 'Times New Roman'
                                run.font.size = Pt(12)
                    for ri, row_data in enumerate(data_rows):
                        for j, cell_text in enumerate(row_data):
                            cell = table.rows[ri + 1].cells[j]
                            cell.text = cell_text
                            for p in cell.paragraphs:
                                for run in p.runs:
                                    run.font.name = 'Times New Roman'
                                    run.font.size = Pt(12)
                    doc.add_paragraph()
            continue

        p = doc.add_paragraph()
        p.add_run(line)
        i += 1

    doc.save(docx_path)
    print(f'Converted: {md_path} -> {docx_path}')

if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('Usage: python convert_md_to_docx.py input.md output.docx')
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
