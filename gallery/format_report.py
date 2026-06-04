import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def format_report(input_path, output_path):
    doc = Document(input_path)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.line_spacing = 1.5
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.first_line_indent = Cm(1.25)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    rFonts.set(qn('w:cs'), 'Times New Roman')
    for p in doc.paragraphs:
        text = p.text.strip().upper()
        for run in p.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 0, 0)
        if text in ('ВВЕДЕНИЕ', 'ЗАКЛЮЧЕНИЕ', 'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.bold = True
        elif text.startswith('1. АНАЛИЗ') or text.startswith('2. РАБОЧЕЕ'):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(1.25)
            for run in p.runs:
                run.bold = True
        elif text and text[0].isdigit() and '. ' in p.text[:5]:
            for run in p.runs:
                run.bold = True
        elif p.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.bold = False
        else:
            for run in p.runs:
                run.bold = False
                run.italic = False
                run.underline = False
    for p in doc.paragraphs:
        if p.text.strip() == '':
            p.paragraph_format.first_line_indent = Cm(0)
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(14)
    doc.save(output_path)
    print(f'Formatted: {input_path} -> {output_path}')

if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('Usage: python format_report.py input.docx output.docx')
        sys.exit(1)
    format_report(sys.argv[1], sys.argv[2])
